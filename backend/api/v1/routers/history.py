"""History and analytics API routes for v1 endpoints."""
from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, List, Optional, Iterable

from fastapi import APIRouter, Depends, HTTPException, Response, status
from fastapi.responses import StreamingResponse

from backend.api.schemas import (
    AggregatedResultsSchema,
    DashboardMetricsResponseSchema,
    JobHistoryDetailSchema,
    JobHistoryListResponseSchema,
    JobHistorySummarySchema,
    JobStatusSchema,
    JobSummaryMetricsSchema,
    LowConfidenceFieldSchema,
    PageExtractionSchema,
    SaveEditsRequestSchema,
    SaveEditsResponseSchema,
    TimeWindowMetricsSchema,
)
from backend.api.schemas.common_schemas import BoundingBoxSchema
from backend.api.schemas.job_schemas import (
    FieldExtractionSchema,
    TableCellSchema,
    TableColumnSchema,
    TableExtractionSchema,
)
from backend.api.v1.dependencies import (
    get_delete_job_handler,
    get_history_job_detail_handler,
    get_history_metrics_handler,
    get_list_history_jobs_handler,
    get_low_confidence_fields_handler,
    get_page_repository,
    get_save_edits_handler,
)
from backend.application.commands.delete_job import DeleteJobCommand, DeleteJobHandler
from backend.application.commands.save_edits import (
    FieldEdit,
    SaveEditsCommand,
    SaveEditsHandler,
    TableCellEdit,
)
from backend.application.dto.history_dto import (
    DashboardMetricsDTO,
    HistoryJobSummaryDTO,
    LowConfidenceFieldDTO,
    TimeWindowMetricsDTO,
)
from backend.application.queries.list_history_jobs import (
    ListHistoryJobsHandler,
    ListHistoryJobsQuery,
)
from backend.application.queries.get_history_job_detail import (
    GetHistoryJobDetailHandler,
    GetHistoryJobDetailQuery,
)
from backend.application.queries.get_history_metrics import (
    GetHistoryMetricsHandler,
    GetHistoryMetricsQuery,
)
from backend.application.queries.list_low_confidence_fields import (
    ListLowConfidenceFieldsHandler,
    ListLowConfidenceFieldsQuery,
)
from backend.repositories.snapshot_repository import load_snapshot, load_canonical
from backend.domain.value_objects import CanonicalFieldIndex
from backend.domain.entities.field_extraction import FieldExtraction
from backend.domain.entities.page_extraction import PageExtraction
from backend.domain.entities.table_extraction import TableExtraction
from backend.domain.exceptions import (
    EntityNotFoundError,
    EntityValidationError,
    RepositoryError,
)
from backend.domain.repositories.page_repository import PageRepository
from backend.models.job import (
    BoundingBox as LegacyBoundingBox,
    FieldExtraction as LegacyFieldExtraction,
    PageExtraction as LegacyPageExtraction,
    TableCell as LegacyTableCell,
    TableColumn as LegacyTableColumn,
    TableExtraction as LegacyTableExtraction,
    ExtractionJob,
)
from openpyxl import Workbook
import re

router = APIRouter(prefix="/history", tags=["history"])


def _build_job_summary(job: ExtractionJob) -> JobSummaryMetricsSchema:
    total_fields = sum(len(page.fields) for page in job.pages)
    total_tables = sum(len(page.tables) for page in job.pages)
    processing_ms: Optional[int] = None
    started_at = job.status.started_at
    finished_at = job.status.finished_at
    if started_at and finished_at:
        processing_ms = int((finished_at - started_at).total_seconds() * 1000)

    return JobSummaryMetricsSchema(
        totalPages=len(job.pages),
        totalFields=total_fields,
        totalTables=total_tables,
        totalProcessingMs=processing_ms,
        startedAt=started_at,
        finishedAt=finished_at,
    )


@router.get("/jobs", response_model=JobHistoryListResponseSchema)
def list_history_jobs(
    handler: ListHistoryJobsHandler = Depends(get_list_history_jobs_handler),
) -> JobHistoryListResponseSchema:
    summaries = handler.handle(ListHistoryJobsQuery())
    jobs = [_history_summary_to_schema(summary) for summary in summaries]
    return JobHistoryListResponseSchema(jobs=jobs)


@router.get("/jobs/{job_id}/canonical.xlsx")
def download_canonical_excel(
    job_id: str,
    handler: GetHistoryJobDetailHandler = Depends(get_history_job_detail_handler),
):
    canonical: Dict[str, Any] | None = None
    try:
        detail = handler.handle(GetHistoryJobDetailQuery(job_id=job_id))
        if detail and detail.canonical:
            canonical = detail.canonical  # type: ignore[assignment]
    except Exception:
        canonical = None

    if canonical is None:
        canonical = load_canonical(job_id)
    if canonical is None:
        snapshot = load_snapshot(job_id)
        if snapshot is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Snapshot not found")
        canonical = snapshot.get("canonical") or {}

    policy = canonical.get("policyConversion") or canonical.get("policy_conversion") or {}
    if not isinstance(policy, dict):
        policy = {}

    # Build a lightweight fallback from snapshot fields if canonical is sparse
    snapshot = load_snapshot(job_id)
    field_lookup: Dict[str, Dict[str, Any]] = {}
    if snapshot and isinstance(snapshot, dict):
        pages = snapshot.get("pages") or []
        for page in pages:
            if not isinstance(page, dict):
                continue
            page_num = page.get("pageNumber") or page.get("page_number")
            for field in page.get("fields", []):
                if not isinstance(field, dict):
                    continue
                name = (field.get("fieldName") or field.get("name") or "").strip()
                if not name:
                    continue
                norm = _normalize_label(name)
                if norm in field_lookup:
                    continue
                field_lookup[norm] = {
                    "value": field.get("value"),
                    "confidence": field.get("confidence"),
                    "sources": [{"page": page_num, "fieldId": field.get("id")}],
                }

    wb = Workbook()
    ws = wb.active
    ws.title = "Policy Conversion"
    ws.append(["Field", "Value", "Confidence", "Source pages"])

    # Always emit all known fields, even if null
    for field in CanonicalFieldIndex.ordered():
        entry = policy.get(field.label)
        value = entry
        confidence = None
        pages_str = ""
        if isinstance(entry, dict):
            value = entry.get("value")
            confidence = entry.get("confidence")
            sources = entry.get("sources") or []
            pages = sorted({source.get("page") for source in sources if isinstance(source, dict) and source.get("page") is not None})
            pages_str = ", ".join(str(page) for page in pages) if pages else ""
        if (value is None or value == ""):
            # Fallback to snapshot field match
            fallback = field_lookup.get(_normalize_label(field.label))
            if fallback:
                value = fallback.get("value") or "not provided"
                confidence = confidence or fallback.get("confidence")
                sources = fallback.get("sources") or []
                pages = sorted({source.get("page") for source in sources if isinstance(source, dict) and source.get("page") is not None})
                pages_str = ", ".join(str(page) for page in pages) if pages else pages_str
        if value is None or value == "":
            value = "not provided"
        ws.append([field.label, value, confidence, pages_str])

    if len(ws["A"]) <= 1:  # only header row
        ws.append(["(no canonical data)", None, None, None])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="canonical_{job_id}.xlsx"'}
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.get("/jobs/{job_id}", response_model=JobHistoryDetailSchema)
def get_history_job(
    job_id: str,
    handler: GetHistoryJobDetailHandler = Depends(get_history_job_detail_handler),
) -> JobHistoryDetailSchema:
    try:
        job = handler.handle(GetHistoryJobDetailQuery(job_id=job_id))
    except EntityNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc

    summary = _build_job_summary(job)
    pages: List[PageExtractionSchema] = []
    for page in job.pages:
        image_url: Optional[str] = None
        if page.image_path and page.image_path.exists():
            image_url = f"/api/jobs/{job.status.job_id}/pages/{page.page_number}/image"
        pages.append(_legacy_page_to_schema(page, image_url))

    aggregated_payload = job.aggregated or {"jobId": job.status.job_id, "fields": []}
    aggregated = AggregatedResultsSchema.model_validate(aggregated_payload)
    document_types = job.canonical.get("documentTypes") if job.canonical else []
    if not isinstance(document_types, list):
        document_types = []
    page_classifications = job.metadata.get("pageClassifications")
    if not isinstance(page_classifications, list):
        page_classifications = []
    status = JobStatusSchema(
        jobId=job.status.job_id,
        totalPages=job.status.total_pages,
        processedPages=job.status.processed_pages,
        state=job.status.state,
        errors=job.status.errors,
        startedAt=job.status.started_at,
        finishedAt=job.status.finished_at,
        documentType=job.document_type,
        documentTypes=document_types or [],
    )

    return JobHistoryDetailSchema(
        jobId=job.status.job_id,
        documentName=job.metadata.get("originalFilename", job.pdf_path.name),
        summary=summary,
        status=status,
        pages=pages,
        aggregated=aggregated,
        metadata=job.metadata,
        documentType=job.document_type,
        canonical=job.canonical,
        mappingTrace=job.mapping_trace or None,
        pageClassifications=page_classifications,
    )


@router.get("/metrics", response_model=DashboardMetricsResponseSchema)
def get_history_metrics(
    handler: GetHistoryMetricsHandler = Depends(get_history_metrics_handler),
) -> DashboardMetricsResponseSchema:
    metrics = handler.handle(GetHistoryMetricsQuery())
    return _dashboard_metrics_to_schema(metrics)


@router.get("/low-confidence", response_model=List[LowConfidenceFieldSchema])
def list_low_confidence_fields(
    limit: int = 50,
    jobId: Optional[str] = None,
    handler: ListLowConfidenceFieldsHandler = Depends(get_low_confidence_fields_handler),
) -> List[LowConfidenceFieldSchema]:
    records = handler.handle(ListLowConfidenceFieldsQuery(limit=limit, job_id=jobId))
    return [
        _low_confidence_field_to_schema(record)
        for record in records
    ]


@router.delete(
    "/jobs/{job_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    response_class=Response,
    response_model=None,
)
def delete_history_job(
    job_id: str,
    handler: DeleteJobHandler = Depends(get_delete_job_handler),
) -> None:
    try:
        handler.handle(DeleteJobCommand(job_id=job_id))
    except EntityNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except EntityValidationError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except RepositoryError as exc:
        raise HTTPException(status_code=500, detail="Failed to delete job") from exc
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/jobs/{job_id}/edits", response_model=SaveEditsResponseSchema)
def save_history_page_edits(
    job_id: str,
    request: SaveEditsRequestSchema,
    handler: SaveEditsHandler = Depends(get_save_edits_handler),
    page_repository: PageRepository = Depends(get_page_repository),
) -> SaveEditsResponseSchema:
    field_edits = [
        FieldEdit(
            page_number=request.page,
            field_name=field.name,
            new_value=field.value,
        )
        for field in request.fields
    ]
    table_cell_edits = [
        TableCellEdit(
            page_number=request.page,
            row=cell.row,
            column=cell.column,
            new_value=cell.value,
        )
        for cell in request.tableCells
    ]

    command = SaveEditsCommand(
        job_id=job_id,
        field_edits=field_edits,
        table_cell_edits=table_cell_edits,
    )

    try:
        handler.handle(command)
    except EntityNotFoundError as exc:
        message = str(exc)
        status_code = 404 if "not found" in message.lower() else 400
        raise HTTPException(status_code=status_code, detail=message) from exc
    except EntityValidationError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except RepositoryError as exc:
        raise HTTPException(status_code=500, detail="Failed to save edits") from exc

    page = page_repository.find_page(job_id, request.page)
    if not page:
        raise HTTPException(status_code=404, detail="Page not found")

    return SaveEditsResponseSchema(
        jobId=job_id,
        page=request.page,
        updatedFields=[_field_to_schema(field) for field in page.fields],
        updatedTables=[_table_to_schema(page, table) for table in page.tables],
    )


def _field_to_schema(field: FieldExtraction) -> FieldExtractionSchema:
    bbox = None
    if field.bounding_box:
        bbox = BoundingBoxSchema(**field.bounding_box.to_dict())

    return FieldExtractionSchema(
        id=str(field.id),
        page=field.page_number,
        name=field.field_name,
        value=field.value,
        confidence=field.confidence.value,
        bbox=bbox,
        sourceType=field.source,
        revised=field.was_edited,
        originalValue=field.normalized_value if field.normalized_value is not None else None,
    )


def _table_to_schema(page: PageExtraction, table: TableExtraction) -> TableExtractionSchema:
    columns: List[TableColumnSchema] = []
    rows: List[List[TableCellSchema]] = []
    data_row_indices: List[int] = []
    has_header_row = False

    if table.num_rows > 0 and table.num_columns > 0:
        cell_map = {(cell.row, cell.column): cell for cell in table.cells}
        header_rows = {cell.row for cell in table.cells if cell.is_header}
        header_row_index = min(header_rows) if header_rows else None
        has_header_row = bool(header_rows)

        for column_index in range(table.num_columns):
            header_cell = cell_map.get((header_row_index, column_index)) if header_row_index is not None else None
            header_value = (header_cell.content.strip() if header_cell and header_cell.content else "")
            columns.append(
                TableColumnSchema(
                    key=f"col_{column_index}",
                    header=header_value or f"Column {column_index + 1}",
                    confidence=(header_cell.confidence.value if header_cell and header_cell.confidence else None),
                )
            )

        data_row_indices = sorted({cell.row for cell in table.cells if cell.row not in header_rows})
        if not data_row_indices:
            data_row_indices = sorted({cell.row for cell in table.cells})

        for row_index in data_row_indices:
            row_cells: List[TableCellSchema] = []
            for column_index in range(table.num_columns):
                cell = cell_map.get((row_index, column_index))
                bbox = (
                    BoundingBoxSchema(**cell.bounding_box.to_dict())
                    if cell and cell.bounding_box
                    else None
                )
                row_cells.append(
                    TableCellSchema(
                        value=cell.content if cell else "",
                        confidence=cell.confidence.value if cell and cell.confidence else None,
                        bbox=bbox,
                    )
                )
            rows.append(row_cells)
    else:
        for column_index in range(max(table.num_columns, 0)):
            columns.append(
                TableColumnSchema(
                    key=f"col_{column_index}",
                    header=f"Column {column_index + 1}",
                )
            )

    bbox = BoundingBoxSchema(**table.bounding_box.to_dict()) if table.bounding_box else None
    row_start_index = data_row_indices[0] if data_row_indices else 0

    return TableExtractionSchema(
        id=str(table.id),
        page=page.page_number,
        caption=table.title,
        confidence=table.confidence.value,
        columns=columns,
        rows=rows,
        bbox=bbox,
        normalized=True,
        tableGroupId=None,
        continuationOf=None,
        inferredHeaders=has_header_row,
        rowStartIndex=row_start_index,
    )



def _history_summary_to_schema(dto: HistoryJobSummaryDTO) -> JobHistorySummarySchema:
    return JobHistorySummarySchema(
        jobId=dto.job_id,
        documentName=dto.document_name,
        documentType=dto.document_type,
        totalPages=dto.total_pages,
        totalFields=dto.total_fields,
        totalTables=dto.total_tables,
        totalProcessingMs=dto.total_processing_ms,
        startedAt=dto.started_at,
        finishedAt=dto.finished_at,
        lastModified=dto.last_modified,
        status=dto.status,
        confidenceBuckets=dto.confidence_buckets,
        lowConfidenceCount=dto.low_confidence_count,
    )


def _dashboard_metrics_to_schema(dto: DashboardMetricsDTO) -> DashboardMetricsResponseSchema:
    return DashboardMetricsResponseSchema(
        week=_time_window_to_schema(dto.week),
        month=_time_window_to_schema(dto.month),
        year=_time_window_to_schema(dto.year),
    )


def _time_window_to_schema(dto: TimeWindowMetricsDTO) -> TimeWindowMetricsSchema:
    return TimeWindowMetricsSchema(
        totalJobs=dto.total_jobs,
        totalPages=dto.total_pages,
        totalFields=dto.total_fields,
        totalTables=dto.total_tables,
        totalProcessingMs=dto.total_processing_ms,
    )


def _low_confidence_field_to_schema(dto: LowConfidenceFieldDTO) -> LowConfidenceFieldSchema:
    return LowConfidenceFieldSchema(
        jobId=dto.job_id,
        documentName=dto.document_name,
        page=dto.page,
        name=dto.name,
        value=dto.value,
        confidence=dto.confidence,
    )


def _legacy_page_to_schema(page: LegacyPageExtraction, image_url: Optional[str]) -> PageExtractionSchema:
    fields = [_legacy_field_to_schema(field) for field in page.fields]
    tables = [_legacy_table_to_schema(table) for table in page.tables]

    return PageExtractionSchema(
        pageNumber=page.page_number,
        status=page.status,
        fields=fields,
        tables=tables,
        imageUrl=image_url,
        markdownText=getattr(page, "markdown_text", None),
        errorMessage=page.error_message,
        rotationApplied=getattr(page, "rotation_applied", 0),
        documentTypeHint=getattr(page, "document_type_hint", None),
        documentTypeConfidence=getattr(page, "document_type_confidence", None),
    )


def _legacy_field_to_schema(field: LegacyFieldExtraction) -> FieldExtractionSchema:
    bbox = _legacy_bbox_to_schema(field.bbox)
    confidence = field.confidence if field.confidence is not None else 0.0

    return FieldExtractionSchema(
        id=field.id,
        page=field.page,
        name=field.name,
        value=field.value,
        confidence=confidence,
        bbox=bbox,
        sourceType=field.source_type,
        revised=field.revised,
        originalValue=field.original_value,
    )


def _legacy_table_to_schema(table: LegacyTableExtraction) -> TableExtractionSchema:
    columns = [_legacy_column_to_schema(column) for column in table.columns]
    rows = [[_legacy_cell_to_schema(cell) for cell in row] for row in table.rows]

    return TableExtractionSchema(
        id=table.id,
        page=table.page,
        caption=table.caption,
        confidence=table.confidence,
        columns=columns,
        rows=rows,
        bbox=_legacy_bbox_to_schema(table.bbox),
        normalized=table.normalized,
        tableGroupId=table.table_group_id,
        continuationOf=table.continuation_of,
        inferredHeaders=table.inferred_headers,
        rowStartIndex=table.row_start_index,
    )


def _legacy_column_to_schema(column: LegacyTableColumn) -> TableColumnSchema:
    return TableColumnSchema(
        key=column.key,
        header=column.header,
        type=column.type,
        confidence=column.confidence,
    )


def _legacy_cell_to_schema(cell: LegacyTableCell) -> TableCellSchema:
    return TableCellSchema(
        value=cell.value,
        confidence=cell.confidence,
        bbox=_legacy_bbox_to_schema(cell.bbox),
    )


def _legacy_bbox_to_schema(bbox: Optional[LegacyBoundingBox]) -> Optional[BoundingBoxSchema]:
    if not bbox:
        return None
    return BoundingBoxSchema(
        x=bbox.x,
        y=bbox.y,
        width=bbox.width,
        height=bbox.height,
    )


def _normalize_label(label: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", label.lower().strip())
